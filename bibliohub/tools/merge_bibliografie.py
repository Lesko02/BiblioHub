"""
Merge bibliografie (v1.1).

Confronta due Excel prodotti dallo Scraper (stesse intestazioni di colonna).
Il confronto è sul testo delle celle, non sulla formattazione Excel, e ignora
maiuscole/minuscole. Ordinamento, Periodo ed Edizione fanno parte della chiave,
insieme a Docente, Descrizione, Data e ISBN.

Modi:
  diff    solo le righe presenti in una sola estrazione (Left / Right)
  common  solo le righe presenti in entrambe (In comune)

Left = estrazione precedente, Right = estrazione corrente.
Le righe "In comune" riportano il testo del file corrente.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import pandas as pd

from ..values import as_text

VERSION = "1.1"

MODE_DIFF = "diff"
MODE_COMMON = "common"

ORIGIN_LEFT = "Left"
ORIGIN_RIGHT = "Right"
ORIGIN_COMMON = "In comune"

REQUIRED_COLUMNS = [
    "Insegnamento", "Periodo", "Edizione",
    "Docente", "Descrizione", "Data", "ISBN",
]
COMPARE_COLUMNS = ["Ordinamento", "Periodo", "Edizione", "Docente", "Descrizione", "Data", "ISBN"]
OUTPUT_COLUMNS = [
    "Insegnamento", "Ordinamento", "Periodo", "Edizione",
    "Docente", "Descrizione", "Data", "ISBN", "Origine",
]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
ORIGIN_FILLS = {
    ORIGIN_LEFT: PatternFill(start_color="FDE9D9", end_color="FDE9D9", fill_type="solid"),
    ORIGIN_RIGHT: PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    ORIGIN_COMMON: PatternFill(start_color="D6E3F0", end_color="D6E3F0", fill_type="solid"),
}


def _display(value) -> str:
    return as_text(value)


def _norm(value) -> str:
    return _display(value).casefold()


def _require_columns(df, path):
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(
            "Il file non ha le colonne dello Scraper "
            "(non basta che siano nello stesso numero).\n"
            f"{path}\n"
            f"Colonne mancanti: {', '.join(missing)}"
        )


def _load(path, log):
    log(f"Lettura: {path}")
    df = pd.read_excel(path)
    _require_columns(df, path)
    if "Ordinamento" not in df.columns:
        df["Ordinamento"] = ""
        log(f"Colonna Ordinamento assente, le righe restano senza ordinamento: {path}")
    return df


def _group_rows(df, min_chars, drop_short):
    """Raggruppa per insegnamento (senza distinguere le maiuscole)."""
    groups = {}
    order = []
    dropped = 0
    for _, row in df.iterrows():
        if drop_short and len(_display(row["Descrizione"])) < min_chars:
            dropped += 1
            continue
        key = _norm(row["Insegnamento"])
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(row)
    return order, groups, dropped


def _index_rows(rows):
    """Prima occorrenza di ogni chiave di confronto, nell'ordine del file."""
    indexed = {}
    for row in rows:
        indexed.setdefault(tuple(_norm(row[c]) for c in COMPARE_COLUMNS), row)
    return indexed


def _emit(output_rows, row, course_name, origin):
    output_rows.append({
        "Insegnamento": course_name,
        "Ordinamento": _display(row["Ordinamento"]),
        "Periodo": _display(row["Periodo"]),
        "Edizione": _display(row["Edizione"]),
        "Docente": _display(row["Docente"]),
        "Descrizione": _display(row["Descrizione"]),
        "Data": _display(row["Data"]),
        "ISBN": _display(row["ISBN"]),
        "Origine": origin,
    })


def _write_excel(rows, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Merge"
    ws.append(OUTPUT_COLUMNS)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    origin_col = OUTPUT_COLUMNS.index("Origine") + 1
    for record in rows:
        ws.append([record[name] for name in OUTPUT_COLUMNS])
        for cell in ws[ws.max_row]:
            cell.number_format = "@"
        origin = ws.cell(ws.max_row, origin_col)
        fill = ORIGIN_FILLS.get(origin.value)
        if fill is not None:
            origin.fill = fill
            origin.font = Font(bold=True)

    for col_idx in range(1, len(OUTPUT_COLUMNS) + 1):
        letter = get_column_letter(col_idx)
        longest = max((len(str(cell.value)) for cell in ws[letter] if cell.value is not None), default=0)
        ws.column_dimensions[letter].width = min(longest + 2, 50)

    wb.save(output_path)


def compare_bibliografie(left_path, right_path, output_path, mode=MODE_DIFF,
                         drop_short=True, min_chars=3, log=print):
    """
    Ritorna un riepilogo. Scrive sempre il file, anche se non ci sono righe.
    Left è l'estrazione precedente, Right quella corrente.
    """
    if mode not in (MODE_DIFF, MODE_COMMON):
        raise ValueError(f"Modalità sconosciuta: {mode}")
    min_chars = int(min_chars)
    if min_chars < 0:
        raise ValueError("Il minimo di caratteri non può essere negativo.")

    df_left = _load(left_path, log)
    df_right = _load(right_path, log)

    left_order, left_groups, dropped_left = _group_rows(df_left, min_chars, drop_short)
    right_order, right_groups, dropped_right = _group_rows(df_right, min_chars, drop_short)

    if drop_short:
        log(f"Scartate descrizioni sotto i {min_chars} caratteri: "
            f"{dropped_left} a sinistra, {dropped_right} a destra.")

    course_order = list(left_order)
    for key in right_order:
        if key not in left_groups:
            course_order.append(key)

    output_rows = []
    counts = {ORIGIN_LEFT: 0, ORIGIN_RIGHT: 0, ORIGIN_COMMON: 0}

    log("Confronto del testo in corso (maiuscole e minuscole ignorate)...")

    for course_key in course_order:
        left_rows = left_groups.get(course_key, [])
        right_rows = right_groups.get(course_key, [])
        sample = (right_rows or left_rows)[0]
        course_name = _display(sample["Insegnamento"])

        left_index = _index_rows(left_rows)
        right_index = _index_rows(right_rows)

        if mode == MODE_COMMON:
            for key, row in right_index.items():
                if key in left_index:
                    _emit(output_rows, row, course_name, ORIGIN_COMMON)
                    counts[ORIGIN_COMMON] += 1
            continue

        for key, row in left_index.items():
            if key not in right_index:
                _emit(output_rows, row, course_name, ORIGIN_LEFT)
                counts[ORIGIN_LEFT] += 1
        for key, row in right_index.items():
            if key not in left_index:
                _emit(output_rows, row, course_name, ORIGIN_RIGHT)
                counts[ORIGIN_RIGHT] += 1

    _write_excel(output_rows, output_path)

    if mode == MODE_COMMON:
        log(f"✅ {counts[ORIGIN_COMMON]} testi in comune → {output_path}")
    elif not output_rows:
        log(f"Nessuna differenza oltre le maiuscole/minuscole. File vuoto salvato in {output_path}")
    else:
        log(f"✅ {counts[ORIGIN_LEFT]} Left, {counts[ORIGIN_RIGHT]} Right → {output_path}")

    return {
        "rows": len(output_rows),
        "left": counts[ORIGIN_LEFT],
        "right": counts[ORIGIN_RIGHT],
        "common": counts[ORIGIN_COMMON],
        "dropped_left": dropped_left,
        "dropped_right": dropped_right,
        "output": output_path,
    }
